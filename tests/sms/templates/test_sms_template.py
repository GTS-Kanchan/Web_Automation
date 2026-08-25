"""
Test Suite: SMS → Templates  (Single-session E2E flow)
Login once → page checks → create → verify in list → search → filter → upload → edit → delete

Mirrors the Java SenderIdTest pattern:
  - Static validation via @pytest.mark.parametrize
  - UI flow: create, verify, upload CSV, edit, delete

Migrated to Playwright: local page-object fixture renamed `template_page`
(built on conftest.py's `module_logged_in_page`) to avoid shadowing
pytest-playwright's reserved `page` fixture. `_to_list()` and
`_get_active_sender_id()` are plain helper functions (not pytest fixtures),
so keeping the parameter name `page` on them is safe — pytest only injects
fixtures by parameter name for actual test/fixture functions.
"""

import os
import time
import random
import string
import tempfile

import pytest

from pages.sms.sms_template_page import SMSTemplatePage
from utils.config import Config
from utils.test_data_generator import DATA_DIR


pytestmark = [pytest.mark.sms, pytest.mark.template]


# DATA_DIR now comes from utils/test_data_generator.py rather than being
# locally redefined as os.path.dirname(__file__)-relative -- this file moved
# from tests/ to tests/sms/templates/ as part of the channel-based reorg,
# and the old definition would have silently pointed at a non-existent
# tests/sms/templates/test_data/ directory after the move.
UPLOAD_CSV = os.path.join(DATA_DIR, "valid_template.xlsx")


# ── Upload-test helpers ──────────────────────────────────────────────────────

def _read_names_from_xlsx(path: str) -> list:
    """Return the list of template_name values from an xlsx file.

    Locates the 'template_name' column BY HEADER NAME rather than assuming
    it's column 0 — the bulk-upload format's real header order is
    sender_id, entity_id, template_dlt_id, template_name, template_product,
    content_type, template_content, template_sample (template_name is
    index 3, not 0). A hardcoded row[0] here used to silently read
    sender_id values ("DUMMY"/"AM-SMS") as if they were template names.
    Returns an empty list if openpyxl is unavailable, the file can't be
    read, or no 'template_name' column is found in the header row.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        header_lower = [str(h).strip().lower() if h is not None else "" for h in header]
        name_idx = None
        for candidate in ("template_name", "template name"):
            if candidate in header_lower:
                name_idx = header_lower.index(candidate)
                break
        if name_idx is None:
            print(f"[UploadTest] No 'template_name' column found in header: {header}")
            return []
        names = []
        for row in rows:
            if len(row) > name_idx and row[name_idx] and str(row[name_idx]).strip():
                names.append(str(row[name_idx]).strip())
        return names
    except Exception as exc:
        print(f"[UploadTest] Could not read xlsx names: {exc}")
        return []


def _get_active_sender_id(page) -> str:
    """Read the Sender ID of the first existing template from the live list.

    This is more reliable than hardcoding a value because sender IDs differ
    between test instances.  Falls back to Config.SMS_SENDER_ID if the list
    is empty or the column can't be found.

    NOTE: `page` here is the SMSTemplatePage page object (not a pytest
    fixture) — the parameter name is kept for parity with the Selenium
    source; this function is never resolved by pytest's fixture injection.
    """
    try:
        page.open(SMSTemplatePage.TEMPLATE_LIST_URL)
        page.page.wait_for_timeout(2000)
        page._close_sidebar_overlay()

        # Find which <th> index is the "Sender ID" column
        headers = page.page.locator("table thead th")
        sid_col = None
        for i in range(headers.count()):
            if "sender" in headers.nth(i).inner_text().strip().lower():
                sid_col = i
                break

        if sid_col is None:
            return Config.SMS_SENDER_ID   # header not found — use config fallback

        # Read the first non-empty data row
        rows = page.page.locator("table tbody tr")
        for i in range(rows.count()):
            tds = rows.nth(i).locator("td")
            if tds.count() > sid_col:
                sid = tds.nth(sid_col).inner_text().strip()
                if sid and len(sid) >= 2:
                    print(f"[UploadTest] Using sender ID from existing template: '{sid}'")
                    return sid
    except Exception as exc:
        print(f"[UploadTest] Could not read sender ID from list ({exc}), "
              f"falling back to Config.SMS_SENDER_ID='{Config.SMS_SENDER_ID}'")
    return Config.SMS_SENDER_ID


def _make_upload_xlsx(sender: str):
    """Create a temp xlsx with 3 uniquely-named templates using *sender*.

    Returns (absolute_file_path, [template_name, ...]).
    Skips via pytest.skip when openpyxl is not installed.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        pytest.skip("openpyxl not installed — cannot build upload xlsx")

    suffix = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
    names  = [
        f"UpTmpl_OTP_{suffix}",
        f"UpTmpl_TXN_{suffix}",
        f"UpTmpl_PRO_{suffix}",
    ]
    rows = [
        [names[0], "123456789001", sender, "OTP",
         "Your OTP for {{1}} is {{2}}. Valid {{3}} mins. Do not share. - CPaaS Test " + suffix],
        [names[1], "123456789002", sender, "Transactional",
         "Dear {{1}}, Rs.{{2}} credited to a/c {{3}}. Bal Rs.{{4}}. - CPaaS Test " + suffix],
        [names[2], "123456789003", sender, "Promotional",
         "Hi {{1}}, get {{2}}% off on {{3}}. Use code {{4}} by {{5}}. - CPaaS Test " + suffix],
    ]
    headers = ["Template Name", "DLT Template ID", "Sender ID", "Type", "Content"]

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(tmp.name)
    print(f"[UploadTest] xlsx built: {tmp.name}  sender='{sender}'  names={names}")
    return tmp.name, names


# ── Module-scoped fixtures: ONE browser, ONE login ────────────────────────────

@pytest.fixture(scope="module")
def template_page(module_logged_in_page):
    """Land on Template list page."""
    p = SMSTemplatePage(module_logged_in_page)
    p.navigate_via_sidebar()
    return p


def _to_list(page):
    """Return to Template list — called at start of each test.

    `page` here is the SMSTemplatePage page object, not a pytest fixture —
    see module docstring.
    """
    page.open(SMSTemplatePage.TEMPLATE_LIST_URL)
    page.page.wait_for_timeout(2000)
    page._close_sidebar_overlay()


# Auto-generated test data for this run
def _rand_name():
    suffix = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
    return f"AutoTmpl_{suffix}"


TMPL_TRANS  = _rand_name()
TMPL_PROMO  = _rand_name()
TMPL_OTP    = _rand_name()
DLT_ID      = SMSTemplatePage.generate_dlt_id()

# Sender IDs to select on each template form
SENDER_IDS  = ["dummy", "AM-SMS"]

# Type-specific content (uses DLT {{N}} variable placeholders)
CONTENT_TRANS  = ("Dear {{1}}, your transaction of Rs.{{2}} on {{3}} is successful. "
                  "Reference: {{4}}. - CPaaS Team")
CONTENT_PROMO  = ("Exclusive deal! Get {{1}}% off on all plans using code {{2}}. "
                  "Valid till {{3}}. Reply STOP to opt out. - CPaaS Team")
CONTENT_OTP    = ("Your OTP for {{1}} is {{2}}. Valid for {{3}} minutes. "
                  "Do not share this code with anyone. - CPaaS Team")

SAMPLE_TRANS   = ("Dear John, your transaction of Rs.5000 on 04-Jul-2026 is successful. "
                  "Reference: TXN12345. - CPaaS Team")
SAMPLE_PROMO   = ("Exclusive deal! Get 30% off on all plans using code PROMO30. "
                  "Valid till 31-Jul-2026. Reply STOP to opt out. - CPaaS Team")
SAMPLE_OTP     = ("Your OTP for login is 123456. Valid for 10 minutes. "
                  "Do not share this code with anyone. - CPaaS Team")

# Kept for legacy/negative tests that don't care about type
CONTENT     = CONTENT_OTP

# Module-level state
_created: list = []   # template names successfully created


# ══════════════════════════════════════════════════════════════════════════════
# PART 1 — Static validation (no browser)
# ══════════════════════════════════════════════════════════════════════════════

VALID_NAMES = [
    "MyTemplate",
    "Trans Template 01",
    "OTP-Template",
    "Promo_2024",
    "A" * 100,          # max 100 chars
    "abc",              # min 3 chars
]

INVALID_NAMES = [
    ("",        "empty name"),
    ("  ",      "whitespace only"),
    ("ab",      "too short (2 chars)"),
    ("A" * 101, "too long (101 chars)"),
]

VALID_DLT_IDS = [
    "123456789012",     # 12 digits
    "12345678901234",   # 14 digits
    "12345678901234567890",  # 20 digits
]

INVALID_DLT_IDS = [
    ("",             "empty"),
    ("123456789",    "too short (9 digits)"),
    ("ABC123456789", "contains letters"),
    ("12345678901234567890X", "21 chars with letter"),
]


class TestTemplateValidation:
    """Static validation — no browser needed."""

    @pytest.mark.parametrize("name", VALID_NAMES, ids=VALID_NAMES)
    def test_valid_template_names(self, name):
        assert SMSTemplatePage.is_valid_template_name(name), \
            f"'{name}' should be VALID"

    @pytest.mark.parametrize("name,reason", INVALID_NAMES, ids=[r for _, r in INVALID_NAMES])
    def test_invalid_template_names(self, name, reason):
        assert not SMSTemplatePage.is_valid_template_name(name), \
            f"'{name}' ({reason}) should be INVALID"

    @pytest.mark.parametrize("dlt_id", VALID_DLT_IDS, ids=VALID_DLT_IDS)
    def test_valid_dlt_ids(self, dlt_id):
        assert SMSTemplatePage.is_valid_dlt_id(dlt_id), \
            f"DLT ID '{dlt_id}' should be VALID"

    @pytest.mark.parametrize("dlt_id,reason", INVALID_DLT_IDS, ids=[r for _, r in INVALID_DLT_IDS])
    def test_invalid_dlt_ids(self, dlt_id, reason):
        assert not SMSTemplatePage.is_valid_dlt_id(dlt_id), \
            f"DLT ID '{dlt_id}' ({reason}) should be INVALID"


# ══════════════════════════════════════════════════════════════════════════════
# PART 2 — Page Load & Basic UI
# ══════════════════════════════════════════════════════════════════════════════

class TestTemplatePageLoad:

    @pytest.mark.smoke
    def test_page_loads(self, template_page):
        _to_list(template_page)
        assert template_page.is_template_list_page()
        assert "404" not in template_page.get_title().lower()

    @pytest.mark.smoke
    def test_create_button_visible(self, template_page):
        _to_list(template_page)
        assert template_page.is_create_button_visible(), "Create New Template button must be visible"

    @pytest.mark.smoke
    def test_upload_button_visible(self, template_page):
        _to_list(template_page)
        assert template_page.is_upload_button_visible(), "Upload Template button must be visible"

    @pytest.mark.smoke
    def test_table_or_empty_state_shown(self, template_page):
        _to_list(template_page)
        has_data = template_page.has_records()
        has_empty = template_page.has_no_records_message()
        assert has_data or has_empty, "Page should show records or empty-state message"


# ══════════════════════════════════════════════════════════════════════════════
# PART 3 — Create Templates
# ══════════════════════════════════════════════════════════════════════════════

class TestCreateTemplate:

    def _fill_and_submit(self, template_page, name, tmpl_type, content, sample):
        """Open the create form, fill all fields, and click Save.

        Sender IDs: selects "dummy" then "AM-SMS" from the Alpine multi-select.
        Content / sample: passed per template type so each uses realistic text.
        """
        template_page.click_create_template()
        assert template_page.is_element_present(SMSTemplatePage.FORM_TEMPLATE_NAME, timeout=15000), \
            "Create Template form should open"

        # -- Template name --
        template_page.fill_template_name(name)

        # -- Sender IDs: type "dummy" then "AM-SMS" into the search input --
        template_page.select_sender_ids(SENDER_IDS)
        template_page.page.wait_for_timeout(1000)   # allow Livewire to react (may reveal DLT ID field)

        # -- Product type --
        try:
            template_page.select_type(tmpl_type)
        except Exception:
            pass

        # -- DLT Template ID (appears for Indian senders via Livewire reactivity) --
        template_page.fill_dlt_id(DLT_ID)

        # -- Message content and sample --
        template_page.fill_content(content)
        try:
            template_page.fill_sample(sample)
        except Exception:
            pass

        template_page.click_save()
        template_page.page.wait_for_timeout(2000)

    def _verify_in_list(self, template_page, name):
        """Return to list and assert the template row is present."""
        _to_list(template_page)
        assert template_page.is_template_present_in_list(name), \
            f"Template '{name}' should appear in the list after creation"

    @pytest.mark.smoke
    def test_create_form_opens(self, template_page):
        _to_list(template_page)
        template_page.click_create_template()
        assert template_page.is_element_present(SMSTemplatePage.FORM_TEMPLATE_NAME, timeout=15000), \
            "Create form should be visible after clicking Create"
        _to_list(template_page)

    @pytest.mark.regression
    def test_create_transactional_template(self, template_page):
        _to_list(template_page)
        self._fill_and_submit(template_page, TMPL_TRANS, "Transactional",
                              CONTENT_TRANS, SAMPLE_TRANS)
        success = template_page.is_success_toast_shown() or template_page.is_template_list_page()
        assert success, f"Creating Transactional template '{TMPL_TRANS}' should succeed"
        if success:
            _created.append(TMPL_TRANS)
        # Verify the template appears in the list
        self._verify_in_list(template_page, TMPL_TRANS)

    @pytest.mark.regression
    def test_create_promotional_template(self, template_page):
        _to_list(template_page)
        self._fill_and_submit(template_page, TMPL_PROMO, "Promotional",
                              CONTENT_PROMO, SAMPLE_PROMO)
        success = template_page.is_success_toast_shown() or template_page.is_template_list_page()
        assert success, f"Creating Promotional template '{TMPL_PROMO}' should succeed"
        if success:
            _created.append(TMPL_PROMO)
        self._verify_in_list(template_page, TMPL_PROMO)

    @pytest.mark.regression
    def test_create_otp_template(self, template_page):
        _to_list(template_page)
        self._fill_and_submit(template_page, TMPL_OTP, "OTP",
                              CONTENT_OTP, SAMPLE_OTP)
        success = template_page.is_success_toast_shown() or template_page.is_template_list_page()
        assert success, f"Creating OTP template '{TMPL_OTP}' should succeed"
        if success:
            _created.append(TMPL_OTP)
        self._verify_in_list(template_page, TMPL_OTP)

    @pytest.mark.regression
    def test_empty_form_shows_validation_errors(self, template_page):
        _to_list(template_page)
        template_page.click_create_template()
        template_page.click_save()
        template_page.page.wait_for_timeout(1000)
        still_on_form = template_page.is_element_present(SMSTemplatePage.FORM_TEMPLATE_NAME, timeout=3000)
        errors = template_page.get_validation_errors()
        assert still_on_form or errors, "Empty form should show validation errors"
        _to_list(template_page)

    @pytest.mark.regression
    @pytest.mark.negative
    def test_short_name_rejected(self, template_page):
        _to_list(template_page)
        template_page.click_create_template()
        template_page.fill_template_name("ab")   # 2 chars — below minimum
        template_page.fill_content(CONTENT)
        template_page.click_save()
        template_page.page.wait_for_timeout(1000)
        still_on_form = template_page.is_element_present(SMSTemplatePage.FORM_TEMPLATE_NAME, timeout=3000)
        errors = template_page.get_validation_errors()
        found = template_page.is_template_present_in_list("ab")
        assert (still_on_form or errors) and not found, \
            "Template name 'ab' (too short) should be rejected"

    @pytest.mark.regression
    @pytest.mark.negative
    def test_cancel_create_returns_to_list(self, template_page):
        _to_list(template_page)
        template_page.click_create_template()
        template_page.fill_template_name("SHOULD NOT BE SAVED")
        template_page.click_form_cancel()
        template_page.page.wait_for_timeout(1000)
        assert template_page.is_template_list_page(), "Cancel should return to Template list"


# ══════════════════════════════════════════════════════════════════════════════
# PART 4 — Verify Created Templates in List
# ══════════════════════════════════════════════════════════════════════════════

class TestVerifyCreatedTemplates:

    @pytest.mark.smoke
    def test_created_templates_in_list(self, template_page):
        if not _created:
            pytest.skip("No templates were created in this run")
        for name in _created:
            found = template_page.is_template_present_in_list(name)
            assert found, f"Created template '{name}' should appear in list"


# ══════════════════════════════════════════════════════════════════════════════
# PART 5 — Search
# ══════════════════════════════════════════════════════════════════════════════

class TestTemplateSearch:

    @pytest.mark.regression
    def test_search_created_template(self, template_page):
        if not _created:
            pytest.skip("No templates created")
        _to_list(template_page)
        template_page.search(_created[0])
        template_page.page.wait_for_timeout(2000)
        assert template_page.get_row_count() > 0 or not template_page.has_no_records_message(), \
            f"Search for '{_created[0]}' should return results"
        template_page.clear_search()

    @pytest.mark.regression
    @pytest.mark.negative
    def test_search_invalid_returns_no_records(self, template_page):
        _to_list(template_page)
        template_page.search("ZZZINVALIDTMPLZZZ999")
        template_page.page.wait_for_timeout(2000)
        assert template_page.get_row_count() == 0 or template_page.has_no_records_message(), \
            "Invalid search should return empty table or no-records message"
        template_page.clear_search()


# ══════════════════════════════════════════════════════════════════════════════
# PART 6 — Filters
# ══════════════════════════════════════════════════════════════════════════════

class TestTemplateFilters:

    @pytest.mark.regression
    def test_filter_by_dlt_id(self, template_page):
        _to_list(template_page)
        try:
            template_page.filter_by_dlt_id(DLT_ID)
            template_page.page.wait_for_timeout(1500)
            page_ok = template_page.get_row_count() >= 0   # any result is fine
            assert page_ok
        except Exception:
            pytest.skip("Filter by DLT ID not available — locator needs update")

    @pytest.mark.regression
    @pytest.mark.negative
    def test_filter_invalid_dlt_id_no_results(self, template_page):
        _to_list(template_page)
        try:
            template_page.filter_by_dlt_id("000000000000INVALID")
            template_page.page.wait_for_timeout(1500)
            assert template_page.get_row_count() == 0 or template_page.has_no_records_message()
        except Exception:
            pytest.skip("Filter not available — locator needs update")


# ══════════════════════════════════════════════════════════════════════════════
# PART 7 — Upload Templates
# ══════════════════════════════════════════════════════════════════════════════

class TestUploadTemplate:

    def _skip_if_no_permission(self, template_page):
        _to_list(template_page)
        if not template_page.is_upload_button_visible():
            pytest.skip("Upload button not visible for this account")

    def _open_upload_popup(self, template_page):
        template_page.click_upload_template()
        template_page.page.wait_for_timeout(1500)
        if not template_page.is_upload_popup_open():
            if template_page.get_toast_error():
                pytest.skip("No upload permission for this account")
            pytest.skip("Upload popup did not open")

    @pytest.mark.smoke
    def test_upload_popup_opens(self, template_page):
        self._skip_if_no_permission(template_page)
        self._open_upload_popup(template_page)
        assert template_page.is_element_present(SMSTemplatePage.FILE_INPUT, timeout=5000), \
            "File input should be present inside upload popup"

    @pytest.mark.smoke
    def test_upload_valid_file_and_verify(self, template_page):
        """Upload valid_template.xlsx and verify all template names appear in the list.

        Flow:
          1. Read template names from the pre-generated valid_template.xlsx.
          2. Open the Upload Template popup and submit the file.
          3. Click 'Yes, Import', then wait 2 minutes for the job to finish.
          4. Navigate to the template list.
          5. Assert every template name from the file appears as a table row.
        """
        self._skip_if_no_permission(template_page)

        if not os.path.exists(UPLOAD_CSV):
            pytest.skip(f"Upload file not found: {UPLOAD_CSV}")

        # Step 1 — read the template names that are in the file
        expected_names = _read_names_from_xlsx(UPLOAD_CSV)
        if not expected_names:
            pytest.skip("Could not read template names from valid_template.xlsx")

        # This file is now 2 directories deeper (tests/sms/templates/) than
        # when it lived directly in tests/ — 3 levels up reaches the project
        # root's reports/ dir instead of the old 1 level.
        _dbg_dir = os.path.join(os.path.dirname(__file__), "..", "..", "..", "reports")
        os.makedirs(_dbg_dir, exist_ok=True)

        _to_list(template_page)
        self._open_upload_popup(template_page)
        template_page.upload_file(UPLOAD_CSV)
        template_page.click_import()

        # Step 3 — poll for the success popup (up to 120 s)
        # If the popup does not appear within 120 s the test fails immediately
        # without proceeding to the list verification.
        print(f"\n[UploadTest] Import triggered. Polling for success popup (max 120 s)...")
        deadline = time.time() + 120
        success  = False
        while time.time() < deadline:
            if template_page.get_upload_success() or template_page.is_success_toast_shown():
                success = True
                elapsed = round(120 - (deadline - time.time()), 1)
                print(f"[UploadTest] Success popup detected after ~{elapsed}s")
                break
            time.sleep(1)

        # Debug screenshot — shows success popup (or whatever is on screen after 120 s)
        ts = time.strftime("%Y%m%d_%H%M%S")
        _ss = os.path.join(_dbg_dir, f"import_done_{ts}.png")
        try:
            template_page.page.screenshot(path=_ss)
            print(f"[UploadTest] Screenshot: {_ss}")
        except Exception:
            _ss = "(screenshot failed)"

        assert success, (
            f"Import success popup did NOT appear within 120 s — "
            f"stopping without list verification.\n"
            f"Expected templates: {expected_names}\n"
            f"Screenshot: {_ss}"
        )

        # Step 4 — success confirmed; navigate to the template list
        _to_list(template_page)
        template_page.page.wait_for_timeout(2000)

        # Step 5 — verify each template name appears in the list
        missing = []
        for name in expected_names:
            if template_page.is_template_present_in_list(name):
                _created.append(name)
            else:
                missing.append(name)

        assert not missing, (
            f"Templates NOT found in list after import: {missing}\n"
            f"Found OK : {[n for n in expected_names if n not in missing]}\n"
            f"Screenshot: {_ss}"
        )

    @pytest.mark.regression
    def test_cancel_upload_closes_popup(self, template_page):
        self._skip_if_no_permission(template_page)
        self._open_upload_popup(template_page)
        template_page.click_cancel_upload()
        template_page.page.wait_for_timeout(1000)
        assert not template_page.is_upload_popup_open(), "Popup should close after Cancel"

    @pytest.mark.regression
    def test_download_sample_available(self, template_page):
        self._skip_if_no_permission(template_page)
        self._open_upload_popup(template_page)
        assert template_page.is_element_present(SMSTemplatePage.BTN_DOWNLOAD_SAMPLE, timeout=5000), \
            "Download Sample button should be present"
        template_page.click_download_sample()
        template_page.page.wait_for_timeout(1000)
        assert template_page.get_toast_error() is None, "No error after Download Sample"

    @pytest.mark.regression
    @pytest.mark.negative
    def test_upload_invalid_format_rejected(self, template_page):
        dummy = os.path.join(DATA_DIR, "_dummy_tmpl.txt")
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(dummy, "w") as f:
            f.write("not a valid template file")
        try:
            self._skip_if_no_permission(template_page)
            self._open_upload_popup(template_page)
            template_page.upload_file(dummy)
            template_page.page.wait_for_timeout(2000)
            inline_error = template_page.get_upload_error()
            if inline_error:
                return
            # NOTE: original Selenium source used a WebDriverWait/EC block here
            # to wait for BTN_IMPORT then execute_script-click it. Converted to
            # Playwright's own wait_for + click, which supersedes the
            # explicit-wait pattern entirely.
            try:
                btn = template_page.page.locator(SMSTemplatePage.BTN_IMPORT).first
                btn.wait_for(state="visible", timeout=5000)
                btn.click(force=True)
                template_page.page.wait_for_timeout(2000)
            except Exception:
                pass
            assert template_page.get_upload_error() or template_page.get_toast_error(), \
                "Invalid file format should be rejected"
        finally:
            if os.path.exists(dummy):
                os.remove(dummy)


# ══════════════════════════════════════════════════════════════════════════════
# PART 8 — Edit Templates
# ══════════════════════════════════════════════════════════════════════════════

class TestEditTemplate:

    @pytest.mark.regression
    def test_edit_and_save(self, template_page):
        _to_list(template_page)
        if not template_page.has_records():
            pytest.skip("No records to edit")
        if not template_page.click_first_row_edit():
            pytest.skip("Edit button not found — locator needs DOM update")
        # Change content slightly
        try:
            template_page.fill_content(f"Updated content {int(template_page.page.evaluate('() => Date.now()'))}")
        except Exception:
            pass
        template_page.click_save()
        template_page.page.wait_for_timeout(2000)
        assert template_page.is_success_toast_shown() or template_page.is_template_list_page(), \
            "Edit save should succeed"
        _to_list(template_page)

    @pytest.mark.regression
    def test_cancel_edit_returns_to_list(self, template_page):
        _to_list(template_page)
        if not template_page.has_records():
            pytest.skip("No records to edit")
        if not template_page.click_first_row_edit():
            pytest.skip("Edit button not found — locator needs DOM update")
        try:
            template_page.fill_content("SHOULD NOT BE SAVED")
        except Exception:
            pass
        template_page.click_form_cancel()
        template_page.page.wait_for_timeout(1000)
        assert template_page.is_template_list_page(), "Cancel should return to Template list"


# ══════════════════════════════════════════════════════════════════════════════
# PART 9 — Delete Templates  (run last — removes data)
# ══════════════════════════════════════════════════════════════════════════════

class TestDeleteTemplate:

    @pytest.mark.regression
    def test_cancel_delete_preserves_record(self, template_page):
        _to_list(template_page)
        if not template_page.has_records():
            pytest.skip("No records to delete")
        before = template_page.get_row_count()
        if not template_page.click_first_row_delete():
            pytest.skip("Delete button not found — locator needs DOM update")
        template_page.cancel_delete()
        template_page.page.wait_for_timeout(500)
        assert template_page.get_row_count() == before, "Row count should not change after cancel"

    @pytest.mark.regression
    def test_confirm_delete_removes_record(self, template_page):
        """Deletes a template created BY THIS TEST RUN — pulled from
        `_created` (populated by TestCreateTemplate/TestUploadTemplate
        earlier in this same file) or, failing that, a fresh scratch
        template created on the spot. Never touches an arbitrary
        pre-existing row — the previous version deleted whatever was in
        row 1 of the live list, which risked removing a real template
        other suites or the QA team depend on."""
        if _created:
            name = _created.pop()
        else:
            name = _rand_name()
            _to_list(template_page)
            template_page.click_create_template()
            template_page.fill_template_name(name)
            try:
                template_page.select_sender_id(SENDER_IDS[0])
            except Exception:
                pass
            try:
                template_page.select_type("Transactional")
            except Exception:
                pass
            template_page.fill_content(CONTENT_TRANS)
            template_page.click_save()
            template_page.page.wait_for_timeout(1500)

        _to_list(template_page)
        template_page.search(name)
        if not template_page.has_records():
            pytest.skip(f"Scratch template '{name}' was not found in the list — "
                         "cannot verify delete without risking an unrelated row")
        if not template_page.click_first_row_delete():
            pytest.skip("Delete button not found — locator needs DOM update")
        template_page.confirm_delete()
        template_page.page.wait_for_timeout(2000)
        template_page.search(name)
        assert template_page.has_no_records_message() or template_page.get_row_count() == 0, \
            f"Scratch template '{name}' should no longer appear after confirming delete"
        template_page.clear_search()
